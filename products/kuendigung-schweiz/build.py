"""Builds "Kündigungsvorlagen Schweiz" – 8 ready-to-fill letters in one Word file.
Run: python products/kuendigung-schweiz/build.py
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

HERE = Path(__file__).parent
DIST = HERE / "dist"
sys.path.insert(0, str(HERE.parent.parent / "tools"))
from cover import make_cover  # noqa: E402

SENDER = "[Vorname Name]\n[Strasse Nr.]\n[PLZ Ort]\n[Telefon / E-Mail]"
RED = RGBColor(0xD5, 0x2B, 0x1E)

# (title, tip, recipient, subject, body)
LETTERS = [
    ("Mietwohnung kündigen (Mieter)",
     "Frist: meist 3 Monate auf einen ortsüblichen Termin (Art. 266c OR) – massgebend ist dein Mietvertrag. "
     "Per Einschreiben senden; der Brief muss vor Fristbeginn beim Vermieter eintreffen. "
     "Bei einer Familienwohnung müssen beide Ehegatten bzw. eingetragenen Partner unterschreiben (Art. 266m OR).",
     "[Vermieter / Verwaltung]\n[Strasse Nr.]\n[PLZ Ort]",
     "Kündigung des Mietvertrags – [Adresse der Wohnung, Stockwerk]",
     "Sehr geehrte Damen und Herren\n\nHiermit kündige ich / kündigen wir den Mietvertrag für die oben genannte Wohnung "
     "fristgerecht per [Kündigungstermin, z. B. 31. März 2027].\n\nBitte bestätigen Sie mir den Erhalt dieser Kündigung "
     "schriftlich und teilen Sie mir einen Termin für die Wohnungsabgabe mit. Für Besichtigungen durch Interessenten bin "
     "ich nach Absprache gerne erreichbar.\n\nDie Mietkaution bitte ich nach der Abgabe auf folgendes Konto freizugeben: "
     "[IBAN, Kontoinhaber]."),
    ("Vorzeitiger Auszug mit Ersatzmieter",
     "Wer ausserhalb von Frist oder Termin auszieht, wird von den Pflichten befreit, wenn er einen zumutbaren, "
     "zahlungsfähigen Ersatzmieter vorschlägt, der den Vertrag zu gleichen Bedingungen übernimmt (Art. 264 OR). "
     "Bewerbungsunterlagen der Interessenten beilegen.",
     "[Vermieter / Verwaltung]\n[Strasse Nr.]\n[PLZ Ort]",
     "Vorzeitige Rückgabe der Wohnung – [Adresse der Wohnung]",
     "Sehr geehrte Damen und Herren\n\nIch möchte die oben genannte Wohnung vorzeitig per [Datum] zurückgeben. "
     "Gemäss Art. 264 OR schlage ich Ihnen folgende zumutbare und zahlungsfähige Nachmieter vor, die bereit sind, "
     "den Mietvertrag zu den gleichen Bedingungen zu übernehmen:\n\n- [Name, Kontakt] – Unterlagen liegen bei\n"
     "- [Name, Kontakt] – Unterlagen liegen bei\n\nIch bitte Sie um eine kurze Rückmeldung bis [Datum]."),
    ("Krankenkasse Grundversicherung kündigen",
     "Ordentliche Franchise ohne alternatives Modell: Kündigung auf Ende Jahr – der Brief muss spätestens am letzten "
     "Arbeitstag im November bei der Kasse eintreffen; auf Ende Juni mit 3 Monaten Frist. Bei Hausarzt-/Telmed-Modellen "
     "und höheren Franchisen gelten oft andere Regeln – Police prüfen. Wechsel erst gültig, wenn die neue Kasse die "
     "Aufnahme bestätigt hat. Per Einschreiben senden.",
     "[Krankenkasse]\n[Strasse Nr.]\n[PLZ Ort]",
     "Kündigung der obligatorischen Grundversicherung – Versicherten-Nr. [Nummer]",
     "Sehr geehrte Damen und Herren\n\nHiermit kündige ich die obligatorische Krankenpflegeversicherung (KVG) für "
     "folgende Person(en) fristgerecht per [31. Dezember 2026]:\n\n- [Name, Geburtsdatum, Versicherten-Nr.]\n\n"
     "Die Bestätigung der neuen Versicherung wird Ihnen von [neue Krankenkasse] zugestellt. "
     "Bitte bestätigen Sie mir die Kündigung schriftlich."),
    ("Krankenkasse Zusatzversicherung kündigen",
     "Zusatzversicherungen (VVG) haben eigene Fristen gemäss Police und AVB, oft 3 Monate auf Ende Jahr. "
     "Tipp: Erst kündigen, wenn eine neue Zusatzversicherung zugesagt hat – bei einer neuen Kasse gibt es "
     "Gesundheitsfragen.",
     "[Krankenkasse]\n[Strasse Nr.]\n[PLZ Ort]",
     "Kündigung der Zusatzversicherung – Police-Nr. [Nummer]",
     "Sehr geehrte Damen und Herren\n\nHiermit kündige ich folgende Zusatzversicherung(en) fristgerecht per "
     "[Datum]:\n\n- [Produktname, versicherte Person]\n\nDie Grundversicherung bleibt davon unberührt. "
     "Bitte bestätigen Sie mir die Kündigung schriftlich."),
    ("Handy- oder Internet-Abo kündigen",
     "Frist und Mindestlaufzeit stehen im Vertrag / in den AGB. Bei Rufnummernmitnahme (Portierung) NICHT selbst "
     "kündigen – das übernimmt der neue Anbieter.",
     "[Anbieter]\n[Strasse Nr.]\n[PLZ Ort]",
     "Kündigung Abonnement – Kundennummer [Nummer], Rufnummer/Anschluss [Nummer]",
     "Sehr geehrte Damen und Herren\n\nHiermit kündige ich mein Abonnement [Name des Abos] fristgerecht "
     "auf den nächstmöglichen Termin, gemäss meiner Berechnung per [Datum].\n\nBitte senden Sie mir eine "
     "schriftliche Bestätigung mit dem genauen Vertragsende und informieren Sie mich, wie allfällige Geräte "
     "zurückzugeben sind."),
    ("Versicherung kündigen (Hausrat, Haftpflicht, Auto)",
     "Frist gemäss Police (oft 3 Monate auf Ablauf). Nach einem Schadenfall besteht in vielen Fällen ein "
     "ausserordentliches Kündigungsrecht (Art. 42 VVG). Bei Prämienerhöhung ist oft eine Kündigung auf den "
     "Zeitpunkt der Erhöhung möglich – AVB prüfen.",
     "[Versicherung]\n[Strasse Nr.]\n[PLZ Ort]",
     "Kündigung Police-Nr. [Nummer] – [Hausrat / Privathaftpflicht / Motorfahrzeug]",
     "Sehr geehrte Damen und Herren\n\nHiermit kündige ich die oben genannte Versicherung fristgerecht per "
     "[Datum] bzw. auf den nächstmöglichen Termin.\n\n[Optional: Die Kündigung erfolgt aufgrund der "
     "Prämienerhöhung vom [Datum] / nach dem Schadenfall Nr. [Nummer].]\n\nBitte bestätigen Sie mir die "
     "Kündigung schriftlich."),
    ("Fitness- oder anderes Abo kündigen",
     "Frist gemäss Vertrag/AGB. Kündigung per Einschreiben, damit du den Zugang beweisen kannst.",
     "[Anbieter]\n[Strasse Nr.]\n[PLZ Ort]",
     "Kündigung Mitgliedschaft – Mitglied-Nr. [Nummer]",
     "Sehr geehrte Damen und Herren\n\nHiermit kündige ich meine Mitgliedschaft fristgerecht per [Datum] bzw. "
     "auf den nächstmöglichen Termin. Bitte stoppen Sie allfällige Lastschriften nach Vertragsende und "
     "bestätigen Sie mir die Kündigung schriftlich."),
    ("Arbeitsstelle kündigen (Arbeitnehmer)",
     "Gesetzliche Fristen ohne abweichende Regelung (Art. 335c OR): 1. Dienstjahr 1 Monat, 2.–9. Dienstjahr "
     "2 Monate, danach 3 Monate – jeweils auf Ende eines Monats; in der Probezeit 7 Tage (Art. 335b OR). "
     "Arbeitsvertrag/GAV gehen vor. Kündigung muss vor Fristbeginn beim Arbeitgeber eintreffen.",
     "[Arbeitgeber]\n[z. Hd. Name, HR]\n[Strasse Nr.]\n[PLZ Ort]",
     "Kündigung meines Arbeitsverhältnisses",
     "Sehr geehrte/r [Frau/Herr Name]\n\nHiermit kündige ich mein Arbeitsverhältnis unter Einhaltung der "
     "vertraglichen Kündigungsfrist per [Datum].\n\nIch danke Ihnen für die gute Zusammenarbeit. Bitte stellen "
     "Sie mir ein qualifiziertes Arbeitszeugnis aus und informieren Sie mich über das Vorgehen bei den "
     "verbleibenden Ferientagen [Anzahl] und der Pensionskasse."),
]


def build_docx(path):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    h = doc.add_heading("Kündigungsvorlagen Schweiz", 0)
    doc.add_paragraph("8 Briefvorlagen zum Ausfüllen – für Word, Pages und Google Docs.")
    doc.add_heading("So gehts", 2)
    for step in (
        "Vorlage wählen, alle [Platzhalter] ersetzen, Tipp-Box löschen.",
        "Datum eintragen, ausdrucken und von Hand unterschreiben.",
        "Wichtige Kündigungen per Einschreiben senden – Quittung aufbewahren.",
        "Fristen im eigenen Vertrag prüfen: Er geht den Tipps in dieser Datei vor.",
    ):
        doc.add_paragraph(step, style="List Number")
    doc.add_paragraph("Hinweis: Vorlagen und Tipps ersetzen keine Rechtsberatung.").runs[0].italic = True
    doc.add_heading("Inhalt", 2)
    for title, *_ in LETTERS:
        doc.add_paragraph(title, style="List Bullet")

    for title, tip, recipient, subject, body in LETTERS:
        doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(title, 1)
        t = doc.add_paragraph()
        r = t.add_run("Tipp: " + tip)
        r.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RED
        doc.add_paragraph("Einschreiben").runs[0].bold = True
        doc.add_paragraph(SENDER)
        doc.add_paragraph(recipient)
        doc.add_paragraph("[Ort], [Datum]")
        doc.add_paragraph(subject).runs[0].bold = True
        for para in body.split("\n\n"):
            doc.add_paragraph(para)
        doc.add_paragraph("Freundliche Grüsse\n\n\n[Unterschrift]\n[Vorname Name]")
    doc.save(path)


if __name__ == "__main__":
    DIST.mkdir(exist_ok=True)
    build_docx(DIST / "Kuendigungsvorlagen-Schweiz.docx")
    make_cover(["Kündigungs-", "vorlagen Schweiz"], "8 Briefe · Word & Google Docs",
               ["Mietwohnung, Ersatzmieter, Krankenkasse (KVG & VVG)",
                "Handy/Internet, Versicherungen, Fitness-Abo, Job",
                "Mit Fristen-Tipps und Gesetzesartikeln (OR, VVG)",
                "Platzhalter ausfüllen, drucken, unterschreiben",
                "Sofort-Download · kein Abo"], DIST / "cover.png")
    print("OK")
