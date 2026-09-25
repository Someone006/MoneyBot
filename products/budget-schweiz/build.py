"""Builds the product files for "Haushaltsbudget Schweiz 2026".

Output (in ./dist): the Excel workbook and a cover image for the listing.
Run: python products/budget-schweiz/build.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).parent
DIST = HERE / "dist"
XLSX = DIST / "Haushaltsbudget-Schweiz-2026.xlsx"
COVER = DIST / "cover.png"

CHF = '"CHF "#,##0.00'
PCT = "0.0%"
RED = "D52B1E"  # Swiss red
DARK = "1F2937"
LIGHT = "F3F4F6"
PALE_RED = "FDECEA"
INPUT = "FFF7D6"

MONTHS = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun", "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]

INCOME = ["Lohn netto", "13. Monatslohn / Bonus", "Nebeneinkommen", "Kinder- / Familienzulagen", "Übrige Einnahmen"]

EXPENSES = {
    "Wohnen": ["Miete / Hypothekarzins", "Nebenkosten", "Strom", "Internet & Handy", "Serafe (Radio/TV)", "Hausrat- & Privathaftpflicht"],
    "Gesundheit": ["Krankenkasse Grundversicherung", "Zusatzversicherungen", "Franchise & Selbstbehalt", "Zahnarzt / Optiker"],
    "Mobilität": ["ÖV (GA / Halbtax / Abo)", "Auto: Benzin / Laden", "Auto: Versicherung & Steuer", "Auto: Leasing / Service", "Parkplatz / Velo"],
    "Steuern": ["Rückstellung Steuern"],
    "Vorsorge & Sparen": ["Säule 3a", "Sparkonto / Notgroschen", "Investieren (ETF etc.)"],
    "Alltag": ["Lebensmittel", "Drogerie & Haushalt", "Kleidung", "Auswärts essen", "Kinder / Betreuung", "Haustiere"],
    "Freizeit": ["Hobbys & Sport", "Ferien & Reisen", "Abos (Streaming, Apps)", "Geschenke & Spenden", "Diverses"],
}
FIXED_GROUPS = ["Wohnen", "Gesundheit", "Mobilität", "Steuern"]

thin = Side(style="thin", color="D1D5DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(size=18, bold=True, color=RED)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(italic=True, color="6B7280")


def build_budget(wb):
    ws = wb.create_sheet("Budget")
    title(ws, "Haushaltsbudget Schweiz 2026", "Gelbe Felder ausfüllen – alles andere rechnet automatisch.")
    headers = ["Kategorie", "Budget / Monat"] + MONTHS + ["Total Jahr", "Ø Monat", "Ø vs. Budget"]
    for col, text in enumerate(headers, 1):
        c = ws.cell(row=4, column=col, value=text)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill(DARK)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "C5"

    first_m, last_m = 3, 14  # C..N
    total_col, avg_col, diff_col = 15, 16, 17
    row = 5
    rows = {"income": [], "groups": {}}

    def section(label):
        nonlocal row
        c = ws.cell(row=row, column=1, value=label)
        c.font = Font(bold=True, color=RED)
        for col in range(1, diff_col + 1):
            ws.cell(row=row, column=col).fill = fill(PALE_RED)
        row += 1

    def item(label, is_expense):
        nonlocal row
        ws.cell(row=row, column=1, value=label).border = BOX
        for col in range(2, last_m + 1):
            c = ws.cell(row=row, column=col)
            c.fill = fill(INPUT)
            c.number_format = CHF
            c.border = BOX
        m1, m2 = get_column_letter(first_m), get_column_letter(last_m)
        ws.cell(row=row, column=total_col, value=f"=SUM({m1}{row}:{m2}{row})")
        ws.cell(row=row, column=avg_col, value=f'=IFERROR(AVERAGE({m1}{row}:{m2}{row}),0)')
        ws.cell(row=row, column=diff_col, value=f"=P{row}-B{row}")
        for col in (total_col, avg_col, diff_col):
            ws.cell(row=row, column=col).number_format = CHF
            ws.cell(row=row, column=col).border = BOX
        row += 1
        return row - 1

    section("EINNAHMEN")
    for label in INCOME:
        rows["income"].append(item(label, False))
    for group, labels in EXPENSES.items():
        section(f"AUSGABEN – {group.upper()}")
        rows["groups"][group] = [item(label, True) for label in labels]
    row += 1

    def total_row(label, formula_for_col, bold=True, fmt=CHF):
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = Font(bold=bold)
        for col in range(2, avg_col + 1):
            c = ws.cell(row=row, column=col, value=formula_for_col(get_column_letter(col)))
            c.number_format = fmt
            c.font = Font(bold=bold)
            c.fill = fill(LIGHT)
            c.border = BOX
        row += 1
        return row - 1

    inc = rows["income"]
    exp = [r for rs in rows["groups"].values() for r in rs]
    save = rows["groups"]["Vorsorge & Sparen"]

    def span(col, rs):
        return f"SUM({col}{rs[0]}:{col}{rs[-1]})"

    def spans(col, groups):
        return "+".join(span(col, rows["groups"][g]) for g in groups)

    t_inc = total_row("Total Einnahmen", lambda c: f"={span(c, inc)}")
    t_exp = total_row("Total Ausgaben (inkl. Sparen)", lambda c: "=" + spans(c, EXPENSES))
    t_bal = total_row("Saldo (Einnahmen – Ausgaben)", lambda c: f"={c}{t_inc}-{c}{t_exp}")
    t_rate = total_row("Sparquote (3a + Sparen + Saldo)",
                       lambda c: f"=IFERROR(({span(c, save)}+{c}{t_bal})/{c}{t_inc},0)", fmt=PCT)
    t_fixed = total_row("davon Fixkosten (Wohnen, Gesundheit, Mobilität, Steuern)",
                        lambda c: "=" + spans(c, FIXED_GROUPS), bold=False)

    last_row = row - 1
    red_font = Font(color="B91C1C", bold=True)
    green_font = Font(color="15803D", bold=True)
    bal_range = f"B{t_bal}:P{t_bal}"
    ws.conditional_formatting.add(bal_range, CellIsRule(operator="lessThan", formula=["0"], font=red_font))
    ws.conditional_formatting.add(bal_range, CellIsRule(operator="greaterThanOrEqual", formula=["0"], font=green_font))
    # Expense lines over budget turn red in the deviation column.
    ws.conditional_formatting.add(f"Q{exp[0]}:Q{exp[-1]}",
                                  CellIsRule(operator="greaterThan", formula=["0.005"], font=red_font))

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 15
    for col in range(first_m, diff_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = 12.5
    ws.column_dimensions[get_column_letter(total_col)].width = 14
    return ws, {"inc": t_inc, "exp": t_exp, "bal": t_bal, "rate": t_rate, "fixed": t_fixed,
                "3a": rows["groups"]["Vorsorge & Sparen"][0], "tax": rows["groups"]["Steuern"][0],
                "last": last_row}


def build_overview(wb, rows):
    ws = wb.create_sheet("Jahresübersicht")
    title(ws, "Jahresübersicht", "Wird automatisch aus dem Blatt «Budget» befüllt.")
    for col, text in enumerate(["Monat", "Einnahmen", "Ausgaben", "Saldo", "Sparquote"], 1):
        c = ws.cell(row=4, column=col, value=text)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill(DARK)
    for i, month in enumerate(MONTHS):
        r = 5 + i
        col = get_column_letter(3 + i)
        ws.cell(row=r, column=1, value=month)
        for j, key in enumerate(["inc", "exp", "bal", "rate"]):
            c = ws.cell(row=r, column=2 + j, value=f"=Budget!{col}{rows[key]}")
            c.number_format = PCT if key == "rate" else CHF
            c.border = BOX
    r = 17
    ws.cell(row=r, column=1, value="Total").font = Font(bold=True)
    for j, col in enumerate("BCD"):
        c = ws.cell(row=r, column=2 + j, value=f"=SUM({col}5:{col}16)")
        c.number_format = CHF
        c.font = Font(bold=True)
    c = ws.cell(row=r, column=5, value=f"=Budget!P{rows['rate']}")
    c.number_format = PCT
    c.font = Font(bold=True)
    for col, width in zip("ABCDE", (12, 16, 16, 16, 12)):
        ws.column_dimensions[col].width = width

    chart = BarChart()
    chart.title = "Einnahmen vs. Ausgaben"
    chart.y_axis.title = "CHF"
    chart.add_data(Reference(ws, min_col=2, max_col=3, min_row=4, max_row=16), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=16))
    chart.height, chart.width = 9, 18
    ws.add_chart(chart, "G4")


def input_row(ws, r, label, value=None, fmt=CHF, note=None):
    ws.cell(row=r, column=1, value=label)
    c = ws.cell(row=r, column=2, value=value)
    c.fill = fill(INPUT)
    c.number_format = fmt
    c.border = BOX
    if note:
        ws.cell(row=r, column=3, value=note).font = Font(italic=True, color="6B7280")


def calc_row(ws, r, label, formula, fmt=CHF, bold=True):
    ws.cell(row=r, column=1, value=label).font = Font(bold=bold)
    c = ws.cell(row=r, column=2, value=formula)
    c.number_format = fmt
    c.font = Font(bold=bold)
    c.fill = fill(LIGHT)
    c.border = BOX


def build_helpers(wb, rows):
    ws = wb.create_sheet("Notgroschen")
    title(ws, "Notgroschen-Rechner", "Faustregel: 3–6 Monate Fixkosten auf der Seite.")
    calc_row(ws, 4, "Fixkosten pro Monat (aus Budget)", f"=Budget!B{rows['fixed']}")
    calc_row(ws, 5, "Ziel minimal (3 Monate)", "=B4*3")
    calc_row(ws, 6, "Ziel komfortabel (6 Monate)", "=B4*6")
    input_row(ws, 7, "Aktuell auf dem Notgroschen-Konto", 0)
    calc_row(ws, 8, "Fortschritt zum 6-Monats-Ziel", "=IFERROR(MIN(B7/B6,1),0)", fmt=PCT)
    calc_row(ws, 9, "Es fehlen noch", "=MAX(B6-B7,0)")
    input_row(ws, 10, "Monatliche Sparrate für den Notgroschen", 0)
    calc_row(ws, 11, "Monate bis zum Ziel", '=IF(B9=0,0,IFERROR(ROUNDUP(B9/B10,0),"Sparrate eingeben"))', fmt="0")

    ws = wb.create_sheet("Säule 3a")
    title(ws, "Säule-3a-Tracker", "Einzahlungen werden aus der Zeile «Säule 3a» im Budget übernommen.")
    input_row(ws, 4, "Maximaler 3a-Betrag für dieses Jahr", None,
              note="Offiziellen Maximalbetrag eintragen (je nach Situation mit/ohne Pensionskasse, siehe ESTV).")
    calc_row(ws, 5, "Bereits eingezahlt", f"=Budget!O{rows['3a']}")
    calc_row(ws, 6, "Noch möglich", "=MAX(B4-B5,0)")
    input_row(ws, 7, "Aktueller Monat (1–12)", 1, fmt="0")
    calc_row(ws, 8, "Nötige Einzahlung pro verbleibendem Monat", "=IFERROR(B6/(13-B7),0)")
    calc_row(ws, 9, "Ausschöpfung", "=IFERROR(B5/B4,0)", fmt=PCT)

    ws = wb.create_sheet("Steuer-Rückstellung")
    title(ws, "Steuer-Rückstellung", "Damit die Steuerrechnung nicht überrascht.")
    input_row(ws, 4, "Erwartete Steuern dieses Jahr", 0,
              note="Tipp: letzte Steuerrechnung oder den Online-Steuerrechner deines Kantons verwenden.")
    calc_row(ws, 5, "Empfohlene Rückstellung pro Monat", "=B4/12")
    calc_row(ws, 6, "Bereits zurückgelegt (aus Budget)", f"=Budget!O{rows['tax']}")
    calc_row(ws, 7, "Noch offen", "=MAX(B4-B6,0)")
    calc_row(ws, 8, "Deckungsgrad", "=IFERROR(B6/B4,0)", fmt=PCT)

    for name in ("Notgroschen", "Säule 3a", "Steuer-Rückstellung"):
        wb[name].column_dimensions["A"].width = 46
        wb[name].column_dimensions["B"].width = 18


def build_start(wb):
    ws = wb.active
    ws.title = "Start"
    title(ws, "Haushaltsbudget Schweiz 2026", "Excel & Google Sheets · in CHF · für Einzelpersonen, Paare und Familien")
    steps = [
        ("So funktioniert's", None),
        ("1.", "Blatt «Budget»: in Spalte B dein geplantes Budget pro Monat eintragen."),
        ("2.", "Jeden Monat die tatsächlichen Beträge in die Monatsspalte eintragen (gelbe Felder)."),
        ("3.", "Saldo, Sparquote und Abweichungen werden automatisch berechnet – rot = über Budget."),
        ("4.", "«Jahresübersicht» zeigt Einnahmen vs. Ausgaben als Diagramm."),
        ("5.", "«Notgroschen», «Säule 3a» und «Steuer-Rückstellung» helfen dir, Reserven aufzubauen."),
        ("", ""),
        ("Tipps", None),
        ("•", "Kategorienamen kannst du frei ändern – die Formeln bleiben korrekt."),
        ("•", "Jährliche Rechnungen (z. B. Serafe, Versicherungen) im Budget durch 12 teilen und monatlich zurücklegen."),
        ("•", "Google Sheets: Datei → Importieren → Hochladen. Alles funktioniert ohne Makros."),
        ("", ""),
        ("Hinweis", None),
        ("", "Diese Vorlage ist ein Planungswerkzeug und keine Steuer- oder Finanzberatung."),
    ]
    r = 4
    for a, b in steps:
        if b is None:
            ws.cell(row=r, column=1, value=a).font = Font(bold=True, size=13, color=DARK)
        else:
            ws.cell(row=r, column=1, value=a).font = Font(bold=True, color=RED)
            ws.cell(row=r, column=2, value=b)
        r += 1
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 100


def build_workbook():
    wb = Workbook()
    build_start(wb)
    _, rows = build_budget(wb)
    build_overview(wb, rows)
    build_helpers(wb, rows)
    wb["Start"].sheet_properties.tabColor = RED
    wb["Budget"].sheet_properties.tabColor = RED
    wb.save(XLSX)


def build_cover():
    w, h = 1600, 1200
    img = Image.new("RGB", (w, h), "#FFFFFF")
    d = ImageDraw.Draw(img)
    bold = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    regular = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    f_big = ImageFont.truetype(bold, 92)
    f_mid = ImageFont.truetype(bold, 50)
    f_small = ImageFont.truetype(regular, 38)
    d.rectangle([0, 0, w, 420], fill="#D52B1E")
    # Swiss cross
    cx, cy, s = 170, 210, 46
    d.rectangle([cx - s * 2, cy - s * 2, cx + s * 2, cy + s * 2], fill="#FFFFFF")
    d.rectangle([cx - s * 0.6, cy - s * 1.6, cx + s * 0.6, cy + s * 1.6], fill="#D52B1E")
    d.rectangle([cx - s * 1.6, cy - s * 0.6, cx + s * 1.6, cy + s * 0.6], fill="#D52B1E")
    d.text((330, 110), "Haushaltsbudget", font=f_big, fill="#FFFFFF")
    d.text((330, 225), "Schweiz 2026", font=f_big, fill="#FFFFFF")
    d.text((80, 480), "Excel & Google Sheets · in CHF", font=f_mid, fill="#1F2937")
    bullets = [
        "Monatsbudget mit Schweizer Kategorien (Krankenkasse, Serafe, 3a …)",
        "Automatische Sparquote, Saldo & Abweichungen",
        "Jahresübersicht mit Diagramm",
        "Notgroschen-, Säule-3a- & Steuer-Rückstellungs-Rechner",
        "Sofort-Download · keine Makros · kein Abo",
    ]
    y = 590
    for b in bullets:
        d.ellipse([90, y + 12, 112, y + 34], fill="#D52B1E")
        d.text((135, y), b, font=f_small, fill="#374151")
        y += 95
    img.save(COVER)


if __name__ == "__main__":
    DIST.mkdir(exist_ok=True)
    build_workbook()
    build_cover()
    print(f"OK: {XLSX}\nOK: {COVER}")
